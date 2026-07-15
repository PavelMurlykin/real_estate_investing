"""Import normalized developer mortgage programs from an XLSX workbook."""

from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from hashlib import sha256
import json
import re
import unicodedata
from zipfile import BadZipFile

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from property.models import CompanyGroup, RealEstateComplex

from .models import Bank, DeveloperMortgageProgram, MortgageProgram
from .mortgage_offer_sync import _normalize_bank_match_name

IMPORT_WORKSHEET_NAME = 'db_import'
MAXIMUM_IMPORT_ROWS = 20_000
DECIMAL_QUANTIZER = Decimal('0.01')
PROGRAM_NAMES_BY_CODE = {
    'market': 'Рыночная ипотека',
    'family': 'Семейная ипотека',
    'it': 'IT-ипотека',
}
SUPPORTED_OBJECT_SCOPES = {'named', 'all_objects'}
REQUIRED_COLUMNS = (
    'record_id',
    'source_sheet_row',
    'company_group_name',
    'real_estate_complex_name',
    'object_scope',
    'bank_name',
    'program_type_code',
    'program_cell_column',
    'program_variant_index',
    'is_program_on_stop',
    'grace_period_rate_percent',
    'grace_period_term_years',
    'interest_rate_percent',
    'initial_payment_min_percent',
    'maximum_loan_term_years',
    'maximum_loan_amount_rub',
    'rate_discount_percent',
    'effective_price_increase_percent',
    'parse_warnings',
    'source_url',
)
MODEL_VALUE_FIELDS = (
    'company_group_id',
    'real_estate_complex_id',
    'bank_id',
    'mortgage_program_id',
    'price_increase_percent',
    'grace_period_months',
    'grace_period_interest_rate',
    'minimum_initial_payment_percent',
    'interest_rate',
    'maximum_loan_term_years',
    'maximum_loan_amount',
    'rate_discount_percent',
    'is_active',
    'source_record_id',
    'source_sheet_row',
    'source_program_column',
    'source_program_variant_index',
)
TARGET_SIGNATURE_FIELDS = MODEL_VALUE_FIELDS[:13]


class DeveloperMortgageProgramImportError(Exception):
    """Raised when the uploaded workbook cannot be safely imported."""


@dataclass(frozen=True)
class DeveloperMortgageProgramImportIssue:
    """Describes one skipped workbook row."""

    workbook_row_number: int
    message: str


@dataclass(frozen=True)
class DeveloperMortgageProgramSourceRecord:
    """Typed normalized record read from the import worksheet."""

    workbook_row_number: int
    source_record_id: int
    source_sheet_row: int
    company_group_name: str
    real_estate_complex_name: str
    object_scope: str
    bank_name: str
    program_type_code: str
    source_program_column: str
    source_program_variant_index: int
    is_program_on_stop: bool
    grace_period_interest_rate: Decimal | None
    grace_period_months: int | None
    interest_rate: Decimal | None
    minimum_initial_payment_percent: Decimal | None
    maximum_loan_term_years: int | None
    maximum_loan_amount: Decimal | None
    rate_discount_percent: Decimal | None
    price_increase_percent: Decimal | None
    source_has_warnings: bool
    source_url: str

    def build_source_key(self):
        """Build a stable key for an expanded source program variant."""
        source_identity = (
            self.source_url,
            self.source_sheet_row,
            normalize_reference_name(self.company_group_name),
            normalize_reference_name(self.real_estate_complex_name),
            self.object_scope,
            _normalize_bank_match_name(self.bank_name),
            self.program_type_code,
            self.source_program_column,
            self.source_program_variant_index,
        )
        serialized_identity = json.dumps(
            source_identity,
            ensure_ascii=False,
            separators=(',', ':'),
        )
        return sha256(serialized_identity.encode('utf-8')).hexdigest()


@dataclass(frozen=True)
class ParsedDeveloperMortgageProgramWorkbook:
    """Contains parsed records and row-level extraction issues."""

    total_rows: int
    records: tuple[DeveloperMortgageProgramSourceRecord, ...]
    issues: tuple[DeveloperMortgageProgramImportIssue, ...]


@dataclass(frozen=True)
class DeveloperMortgageProgramImportResult:
    """Summarizes one completed developer mortgage program import."""

    total_rows: int
    parsed_rows: int
    created: int
    updated: int
    unchanged: int
    skipped: int
    duplicate_rows: int
    source_warning_rows: int
    inactive_rows: int
    issue_messages: tuple[str, ...]

    def to_session_data(self):
        """Return a compact serializable summary for the result page."""
        return {
            'total_rows': self.total_rows,
            'parsed_rows': self.parsed_rows,
            'created': self.created,
            'updated': self.updated,
            'unchanged': self.unchanged,
            'skipped': self.skipped,
            'duplicate_rows': self.duplicate_rows,
            'source_warning_rows': self.source_warning_rows,
            'inactive_rows': self.inactive_rows,
            'issue_messages': list(self.issue_messages),
        }


def normalize_reference_name(value):
    """Normalize a name without merging materially different values."""
    normalized_value = unicodedata.normalize('NFKC', str(value or ''))
    normalized_value = normalized_value.replace('ё', 'е').replace('Ё', 'Е')
    return re.sub(r'\s+', ' ', normalized_value).strip().casefold()


def parse_developer_mortgage_program_workbook(workbook_file):
    """Extract typed records from the normalized XLSX import worksheet."""
    try:
        workbook_file.seek(0)
        workbook = load_workbook(
            workbook_file,
            read_only=True,
            data_only=True,
        )
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        ValueError,
    ) as error:
        raise DeveloperMortgageProgramImportError(
            'Не удалось прочитать XLSX-файл.'
        ) from error

    try:
        if IMPORT_WORKSHEET_NAME not in workbook.sheetnames:
            raise DeveloperMortgageProgramImportError(
                f'В файле отсутствует лист «{IMPORT_WORKSHEET_NAME}».'
            )

        worksheet = workbook[IMPORT_WORKSHEET_NAME]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as error:
            raise DeveloperMortgageProgramImportError(
                'Лист импорта пуст.'
            ) from error

        column_indices = build_column_indices(header_row)
        records = []
        issues = []
        total_rows = 0
        seen_source_record_ids = set()

        for workbook_row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, '') for value in row):
                continue
            total_rows += 1
            if total_rows > MAXIMUM_IMPORT_ROWS:
                raise DeveloperMortgageProgramImportError(
                    'Файл содержит больше '
                    f'{MAXIMUM_IMPORT_ROWS} строк данных.'
                )

            row_values = {
                column_name: (
                    row[column_index]
                    if column_index < len(row)
                    else None
                )
                for column_name, column_index in column_indices.items()
            }
            try:
                record = parse_source_record(
                    row_values,
                    workbook_row_number,
                )
                if record.source_record_id in seen_source_record_ids:
                    raise ValueError(
                        'дублируется record_id '
                        f'{record.source_record_id}'
                    )
                seen_source_record_ids.add(record.source_record_id)
                records.append(record)
            except ValueError as error:
                issues.append(
                    DeveloperMortgageProgramImportIssue(
                        workbook_row_number=workbook_row_number,
                        message=str(error),
                    )
                )

        if not records:
            raise DeveloperMortgageProgramImportError(
                'В файле нет корректных строк для импорта.'
            )

        return ParsedDeveloperMortgageProgramWorkbook(
            total_rows=total_rows,
            records=tuple(records),
            issues=tuple(issues),
        )
    finally:
        workbook.close()


def build_column_indices(header_row):
    """Validate the header and return required column positions."""
    column_indices = {}
    duplicate_columns = set()
    for column_index, raw_column_name in enumerate(header_row):
        column_name = str(raw_column_name or '').strip()
        if column_name not in REQUIRED_COLUMNS:
            continue
        if column_name in column_indices:
            duplicate_columns.add(column_name)
            continue
        column_indices[column_name] = column_index

    if duplicate_columns:
        duplicate_list = ', '.join(sorted(duplicate_columns))
        raise DeveloperMortgageProgramImportError(
            f'В заголовке повторяются столбцы: {duplicate_list}.'
        )

    missing_columns = sorted(set(REQUIRED_COLUMNS) - set(column_indices))
    if missing_columns:
        missing_list = ', '.join(missing_columns)
        raise DeveloperMortgageProgramImportError(
            f'В листе импорта отсутствуют столбцы: {missing_list}.'
        )
    return column_indices


def parse_source_record(row_values, workbook_row_number):
    """Convert one worksheet row to a typed source record."""
    object_scope = parse_required_text(
        row_values['object_scope'],
        'object_scope',
    ).casefold()
    real_estate_complex_name = parse_optional_text(
        row_values['real_estate_complex_name']
    )
    if object_scope == 'named' and not real_estate_complex_name:
        raise ValueError('не указан real_estate_complex_name')

    grace_period_years = parse_optional_decimal(
        row_values['grace_period_term_years'],
        'grace_period_term_years',
    )
    grace_period_months = None
    if grace_period_years is not None:
        grace_period_months_decimal = grace_period_years * Decimal('12')
        if (
            grace_period_months_decimal
            != grace_period_months_decimal.to_integral()
        ):
            raise ValueError(
                'grace_period_term_years не переводится в целое число месяцев'
            )
        grace_period_months = int(grace_period_months_decimal)
        if grace_period_months < 1:
            raise ValueError('grace_period_term_years должен быть больше нуля')

    maximum_loan_term_years = parse_optional_integer(
        row_values['maximum_loan_term_years'],
        'maximum_loan_term_years',
    )
    if maximum_loan_term_years is not None and maximum_loan_term_years < 1:
        raise ValueError('maximum_loan_term_years должен быть больше нуля')

    grace_period_interest_rate = parse_rate(
        row_values['grace_period_rate_percent'],
        'grace_period_rate_percent',
    )
    interest_rate = parse_rate(
        row_values['interest_rate_percent'],
        'interest_rate_percent',
    )
    minimum_initial_payment_percent = parse_rate(
        row_values['initial_payment_min_percent'],
        'initial_payment_min_percent',
    )
    rate_discount_percent = parse_rate(
        row_values['rate_discount_percent'],
        'rate_discount_percent',
    )
    maximum_loan_amount = parse_optional_decimal(
        row_values['maximum_loan_amount_rub'],
        'maximum_loan_amount_rub',
    )
    if maximum_loan_amount is not None and maximum_loan_amount < 0:
        raise ValueError('maximum_loan_amount_rub не может быть отрицательной')

    return DeveloperMortgageProgramSourceRecord(
        workbook_row_number=workbook_row_number,
        source_record_id=parse_positive_integer(
            row_values['record_id'],
            'record_id',
        ),
        source_sheet_row=parse_positive_integer(
            row_values['source_sheet_row'],
            'source_sheet_row',
        ),
        company_group_name=parse_required_text(
            row_values['company_group_name'],
            'company_group_name',
        ),
        real_estate_complex_name=real_estate_complex_name,
        object_scope=object_scope,
        bank_name=parse_required_text(
            row_values['bank_name'],
            'bank_name',
        ),
        program_type_code=parse_required_text(
            row_values['program_type_code'],
            'program_type_code',
        ).casefold(),
        source_program_column=parse_program_column(
            row_values['program_cell_column']
        ),
        source_program_variant_index=parse_positive_integer(
            row_values['program_variant_index'],
            'program_variant_index',
        ),
        is_program_on_stop=parse_boolean(
            row_values['is_program_on_stop'],
            'is_program_on_stop',
        ),
        grace_period_interest_rate=grace_period_interest_rate,
        grace_period_months=grace_period_months,
        interest_rate=interest_rate,
        minimum_initial_payment_percent=(
            minimum_initial_payment_percent
        ),
        maximum_loan_term_years=maximum_loan_term_years,
        maximum_loan_amount=maximum_loan_amount,
        rate_discount_percent=rate_discount_percent,
        price_increase_percent=parse_optional_decimal(
            row_values['effective_price_increase_percent'],
            'effective_price_increase_percent',
        ),
        source_has_warnings=bool(
            parse_optional_text(row_values['parse_warnings'])
        ),
        source_url=parse_optional_text(row_values['source_url']),
    )


def parse_required_text(value, field_name):
    """Return a trimmed required text value."""
    parsed_value = parse_optional_text(value)
    if not parsed_value:
        raise ValueError(f'не заполнен {field_name}')
    return parsed_value


def parse_optional_text(value):
    """Return a trimmed text value or an empty string."""
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


def parse_optional_decimal(value, field_name):
    """Parse a finite optional decimal using spreadsheet-safe conversion."""
    if value in (None, ''):
        return None
    if isinstance(value, bool):
        raise ValueError(f'{field_name} должен быть числом')
    normalized_value = str(value).replace('\xa0', '').replace(' ', '')
    normalized_value = normalized_value.replace(',', '.')
    try:
        parsed_value = Decimal(normalized_value)
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f'{field_name} должен быть числом') from error
    if not parsed_value.is_finite():
        raise ValueError(f'{field_name} должен быть конечным числом')
    return parsed_value.quantize(
        DECIMAL_QUANTIZER,
        rounding=ROUND_HALF_UP,
    )


def parse_rate(value, field_name):
    """Parse an optional percentage constrained to 0–100."""
    parsed_value = parse_optional_decimal(value, field_name)
    if (
        parsed_value is not None
        and not Decimal('0') <= parsed_value <= Decimal('100')
    ):
        raise ValueError(f'{field_name} должен быть от 0 до 100')
    return parsed_value


def parse_optional_integer(value, field_name):
    """Parse an optional integral spreadsheet number."""
    if value in (None, ''):
        return None
    parsed_value = parse_optional_decimal(value, field_name)
    if parsed_value != parsed_value.to_integral():
        raise ValueError(f'{field_name} должен быть целым числом')
    return int(parsed_value)


def parse_positive_integer(value, field_name):
    """Parse a required positive integral spreadsheet number."""
    parsed_value = parse_optional_integer(value, field_name)
    if parsed_value is None or parsed_value < 1:
        raise ValueError(f'{field_name} должен быть положительным числом')
    return parsed_value


def parse_boolean(value, field_name):
    """Parse a normalized boolean cell."""
    if isinstance(value, bool):
        return value
    if value in (0, '0', 'false', 'False', 'нет', 'Нет'):
        return False
    if value in (1, '1', 'true', 'True', 'да', 'Да'):
        return True
    raise ValueError(f'{field_name} должен быть логическим значением')


def parse_program_column(value):
    """Validate the source program column identifier."""
    column_name = parse_required_text(value, 'program_cell_column').upper()
    if column_name not in {'F', 'G', 'H'}:
        raise ValueError('program_cell_column должен быть F, G или H')
    return column_name


def import_developer_mortgage_programs(workbook_file):
    """Parse and load one workbook without creating dictionary records."""
    parsed_workbook = parse_developer_mortgage_program_workbook(workbook_file)
    return load_developer_mortgage_programs(parsed_workbook)


@transaction.atomic
def load_developer_mortgage_programs(parsed_workbook):
    """Resolve references and bulk upsert parsed developer programs."""
    issues = list(parsed_workbook.issues)
    company_group_lookup = build_unique_lookup(
        CompanyGroup.objects.all(),
        lambda company_group: normalize_reference_name(company_group.name),
    )
    bank_lookup = build_unique_lookup(
        Bank.objects.all(),
        lambda bank: _normalize_bank_match_name(bank.name),
    )
    mortgage_program_lookup = build_unique_lookup(
        MortgageProgram.objects.all(),
        lambda mortgage_program: normalize_reference_name(
            mortgage_program.name
        ),
    )
    real_estate_complex_lookup = build_unique_lookup(
        RealEstateComplex.objects.select_related(
            'developer__company_group'
        ).filter(developer__company_group__isnull=False),
        lambda real_estate_complex: (
            real_estate_complex.developer.company_group_id,
            normalize_reference_name(real_estate_complex.name),
        ),
    )

    resolved_records = []
    source_warning_rows = 0
    inactive_rows = 0
    for source_record in parsed_workbook.records:
        if source_record.source_has_warnings:
            source_warning_rows += 1
        if source_record.is_program_on_stop:
            inactive_rows += 1

        resolved_record, issue = resolve_source_record(
            source_record,
            company_group_lookup,
            bank_lookup,
            mortgage_program_lookup,
            real_estate_complex_lookup,
        )
        if issue is not None:
            issues.append(issue)
            continue
        resolved_records.append(resolved_record)

    unique_records = []
    seen_target_signatures = set()
    duplicate_rows = 0
    for resolved_record in resolved_records:
        target_signature = build_target_signature(resolved_record)
        if target_signature in seen_target_signatures:
            duplicate_rows += 1
            continue
        seen_target_signatures.add(target_signature)
        unique_records.append(resolved_record)

    source_keys = [
        source_record.build_source_key()
        for source_record, _model_values in unique_records
    ]
    existing_programs = {
        developer_program.source_key: developer_program
        for developer_program in (
            DeveloperMortgageProgram.objects.filter(
                source_key__in=source_keys
            )
            .order_by()
            .select_for_update()
        )
    }
    programs_to_create = []
    programs_to_update = []
    unchanged = 0
    update_timestamp = timezone.now()

    for source_record, model_values in unique_records:
        source_key = source_record.build_source_key()
        existing_program = existing_programs.get(source_key)
        if existing_program is None:
            programs_to_create.append(
                DeveloperMortgageProgram(
                    source_key=source_key,
                    **model_values,
                )
            )
            continue

        has_changes = False
        for field_name, field_value in model_values.items():
            if getattr(existing_program, field_name) == field_value:
                continue
            setattr(existing_program, field_name, field_value)
            has_changes = True

        if has_changes:
            existing_program.updated_at = update_timestamp
            programs_to_update.append(existing_program)
        else:
            unchanged += 1

    DeveloperMortgageProgram.objects.bulk_create(
        programs_to_create,
        batch_size=500,
    )
    if programs_to_update:
        DeveloperMortgageProgram.objects.bulk_update(
            programs_to_update,
            fields=(*MODEL_VALUE_FIELDS, 'updated_at'),
            batch_size=500,
        )

    skipped = len(issues) + duplicate_rows
    return DeveloperMortgageProgramImportResult(
        total_rows=parsed_workbook.total_rows,
        parsed_rows=len(parsed_workbook.records),
        created=len(programs_to_create),
        updated=len(programs_to_update),
        unchanged=unchanged,
        skipped=skipped,
        duplicate_rows=duplicate_rows,
        source_warning_rows=source_warning_rows,
        inactive_rows=inactive_rows,
        issue_messages=summarize_import_issues(issues),
    )


def build_unique_lookup(queryset, key_builder):
    """Build a lookup that retains ambiguity instead of picking a record."""
    lookup = {}
    for model_object in queryset:
        lookup.setdefault(key_builder(model_object), []).append(model_object)
    return lookup


def resolve_source_record(
    source_record,
    company_group_lookup,
    bank_lookup,
    mortgage_program_lookup,
    real_estate_complex_lookup,
):
    """Resolve one source row exclusively against existing dictionaries."""
    if source_record.object_scope not in SUPPORTED_OBJECT_SCOPES:
        return None, build_issue(
            source_record,
            'область действия '
            f'«{source_record.object_scope}» не поддерживается моделью',
        )

    company_group, issue_message = resolve_unique_lookup_value(
        company_group_lookup,
        normalize_reference_name(source_record.company_group_name),
        'группа компаний',
        source_record.company_group_name,
    )
    if issue_message:
        return None, build_issue(source_record, issue_message)

    bank, issue_message = resolve_unique_lookup_value(
        bank_lookup,
        _normalize_bank_match_name(source_record.bank_name),
        'банк',
        source_record.bank_name,
    )
    if issue_message:
        return None, build_issue(source_record, issue_message)

    mortgage_program_name = PROGRAM_NAMES_BY_CODE.get(
        source_record.program_type_code
    )
    if mortgage_program_name is None:
        return None, build_issue(
            source_record,
            'неизвестный тип ипотечной программы '
            f'«{source_record.program_type_code}»',
        )
    mortgage_program, issue_message = resolve_unique_lookup_value(
        mortgage_program_lookup,
        normalize_reference_name(mortgage_program_name),
        'ипотечная программа',
        mortgage_program_name,
    )
    if issue_message:
        return None, build_issue(source_record, issue_message)

    real_estate_complex = None
    if source_record.object_scope == 'named':
        real_estate_complex, issue_message = resolve_unique_lookup_value(
            real_estate_complex_lookup,
            (
                company_group.pk,
                normalize_reference_name(
                    source_record.real_estate_complex_name
                ),
            ),
            'ЖК',
            source_record.real_estate_complex_name,
        )
        if issue_message:
            return None, build_issue(source_record, issue_message)

    model_values = {
        'company_group_id': company_group.pk,
        'real_estate_complex_id': (
            real_estate_complex.pk if real_estate_complex else None
        ),
        'bank_id': bank.pk,
        'mortgage_program_id': mortgage_program.pk,
        'price_increase_percent': source_record.price_increase_percent,
        'grace_period_months': source_record.grace_period_months,
        'grace_period_interest_rate': (
            source_record.grace_period_interest_rate
        ),
        'minimum_initial_payment_percent': (
            source_record.minimum_initial_payment_percent
        ),
        'interest_rate': source_record.interest_rate,
        'maximum_loan_term_years': (
            source_record.maximum_loan_term_years
        ),
        'maximum_loan_amount': source_record.maximum_loan_amount,
        'rate_discount_percent': source_record.rate_discount_percent,
        'is_active': not source_record.is_program_on_stop,
        'source_record_id': source_record.source_record_id,
        'source_sheet_row': source_record.source_sheet_row,
        'source_program_column': source_record.source_program_column,
        'source_program_variant_index': (
            source_record.source_program_variant_index
        ),
    }
    return (source_record, model_values), None


def resolve_unique_lookup_value(lookup, lookup_key, label, source_value):
    """Resolve one existing dictionary value and report missing ambiguity."""
    matches = lookup.get(lookup_key, [])
    if not matches:
        return None, f'{label} «{source_value}» отсутствует в базе'
    if len(matches) > 1:
        return None, f'{label} «{source_value}» сопоставляется неоднозначно'
    return matches[0], ''


def build_issue(source_record, message):
    """Build a row issue with workbook traceability."""
    return DeveloperMortgageProgramImportIssue(
        workbook_row_number=source_record.workbook_row_number,
        message=message,
    )


def build_target_signature(resolved_record):
    """Return the business projection used to remove exact duplicates."""
    _source_record, model_values = resolved_record
    return tuple(
        model_values[field_name]
        for field_name in TARGET_SIGNATURE_FIELDS
    )


def summarize_import_issues(issues, maximum_messages=30):
    """Aggregate repeated row errors into a compact operator report."""
    issue_counts = Counter(issue.message for issue in issues)
    first_rows = {}
    for issue in issues:
        first_rows.setdefault(issue.message, issue.workbook_row_number)

    ordered_messages = sorted(
        issue_counts,
        key=lambda message: (-issue_counts[message], message),
    )
    summaries = []
    for message in ordered_messages[:maximum_messages]:
        count = issue_counts[message]
        summaries.append(
            f'{message}: {count} строк, пример — строка '
            f'{first_rows[message]} листа {IMPORT_WORKSHEET_NAME}'
        )
    if len(ordered_messages) > maximum_messages:
        summaries.append(
            'Дополнительных типов ошибок: '
            f'{len(ordered_messages) - maximum_messages}.'
        )
    return tuple(summaries)
